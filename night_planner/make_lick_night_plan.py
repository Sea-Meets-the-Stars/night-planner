# Created by JXP and Claude
"""Created by JXP and Claude

Build a Lick Shane-3m + Kast Night-plan workbook, modeled on the
professional night plans in
/mnt/scratch/xavier/Observing/Observations/Lick_Kast/obs_docs/2026
(e.g. Lick_2026A-4_Night_plan.xlsx), but with one science sheet per
target, named exactly by its TNS.

Inputs
------
- possible_targs xlsx (sheet 'possible_targs'): columns
  TNS, RA_HMS, DEC_DMS, Epoch, Pri_mag, Sec_mag (RA-sorted, vetted).
- optional targets CSV: columns including TNS, Pri_POx, FRB_tags
  (merged on TNS to enrich each target's Comments).
- template xlsx (READ-ONLY): its 'Observing Checklist' sheet is copied
  (with 'Done' flags reset) so the pre/post-obs workflow carries over.

Output
------
- An xlsx workbook with sheets, in order:
  'Observing Checklist', one sheet per target (RA order, named by TNS),
  'FRB Observing Summary', 'Post Observing Tracking'.

Exposure times and the Kast red/blue setup are deliberately left blank:
they depend on the observer's grating/dichroic choice and are set at
the telescope.
"""

import argparse
import os

import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

# Timeline header, columns A -> O, exactly as in the template night plans
TIMELINE_HEADER = [
    'Start (LST)', 'Start (Universal Time)', 'Start (PST)',
    'Duration \n(hh:min)', 'End (PST)', 'Action', 'target', 'RA', 'DEC',
    'mag (r)', 'sec mag', 'red side', 'blue side', 'Redshift;', 'Comments',
]

# Column widths lifted from the template timeline sheets (A -> O, then U)
TIMELINE_WIDTHS = {
    'A': 18.63, 'B': 22.0, 'C': 9.63, 'D': 16.5, 'E': 7.38, 'F': 17.13,
    'G': 16.88, 'H': 13.25, 'I': 11.88, 'J': 9.63, 'K': 8.5, 'L': 11.0,
    'M': 12.5, 'N': 14.0, 'O': 44.5, 'Q': 12.0, 'U': 20.25,
}

# Spectrophotometric standard-star options (RA/DEC as text; the template
# stored some of these as Excel time cells, which mangles declinations)
STANDARD_STARS = [
    ('Feige34',    '10:39:36.74', '+43:06:09.3', 2000.0, 11.2),
    ('Feige110',   '23:19:58.39', '-05:09:55.8', 2000.0, 11.8),
    ('BD+28_4211', '21:51:11.02', '+28:51:50.4', 2000.0, 10.8),
    ('BD+33_2642', '15:51:59.9',  '+32:56:54.3', 2000.0, 10.8),
    ('HZ_44',      '13:23:35.26', '+36:07:59.5', 2000.0, 11.7),
    ('G191B2B',    '05:05:30.60', '+52:49:54.0', 2000.0, 11.8),
    ('BD+17_4708', '22:11:31.37', '+18:05:34.2', 2000.0, 9.47),
]

STANDARD_NOTE = (
    'NOTE: Take a spectrophotometric standard at evening & morning '
    'twilight (choose an appropriate August standard, e.g. BD+28 4211 / '
    'Feige 110 / HZ 44 -- pick per Kast setup).'
)

HEADER_FILL = PatternFill(start_color='FFD9D9D9', end_color='FFD9D9D9',
                          fill_type='solid')


def read_targets(possible_targs_path, targets_csv_path=None):
    """Created by JXP and Claude

    Read the vetted target list and (optionally) merge in P(O|x) and
    FRB tags from the targets CSV.

    Inputs
    ------
    possible_targs_path : str
        Path to the possible_targs xlsx (sheet 'possible_targs').
    targets_csv_path : str or None
        Optional path to the targets CSV with Pri_POx and FRB_tags.

    Outputs
    -------
    pandas.DataFrame with columns TNS, RA_HMS, DEC_DMS, Epoch, Pri_mag,
    Sec_mag and (if the CSV was given) Pri_POx, FRB_tags. RA order is
    preserved from the input file.
    """
    targs = pd.read_excel(possible_targs_path, sheet_name='possible_targs')
    if targets_csv_path is not None and os.path.isfile(targets_csv_path):
        extra = pd.read_csv(targets_csv_path)
        keep = [c for c in ['TNS', 'Pri_POx', 'FRB_tags'] if c in extra.columns]
        targs = targs.merge(extra[keep], on='TNS', how='left')
    return targs


def copy_checklist(template_path, wb):
    """Created by JXP and Claude

    Copy the 'Observing Checklist' sheet from the template workbook into
    wb, resetting every 'Done' flag to False (this run has not happened).

    Inputs
    ------
    template_path : str
        Path to the template night-plan xlsx (opened read-only).
    wb : openpyxl.Workbook
        Workbook to receive the new 'Observing Checklist' sheet.

    Outputs
    -------
    The new openpyxl worksheet (also appended to wb).
    """
    twb = openpyxl.load_workbook(template_path, read_only=True,
                                 data_only=False)
    tws = twb['Observing Checklist']

    ws = wb.create_sheet('Observing Checklist')
    for row in tws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            if cell.column_letter == 'B' and cell.row > 1:
                # Reset Done flags for the upcoming run
                if isinstance(cell.value, bool):
                    ws.cell(row=cell.row, column=cell.column, value=False)
                continue
            ws.cell(row=cell.row, column=cell.column, value=cell.value)
    twb.close()

    # Light formatting: header + section rows in bold
    for coord in ('A1', 'B1', 'C1'):
        ws[coord].font = Font(bold=True)
        ws[coord].fill = HEADER_FILL
    for row in ws.iter_rows(min_col=1, max_col=1):
        val = row[0].value
        if val in ('Pre Obs', 'Post Obs') or (
                isinstance(val, str) and val.startswith('Proceed with')):
            row[0].font = Font(bold=True)
    ws.column_dimensions['A'].width = 45.0
    ws.column_dimensions['B'].width = 8.0
    ws.column_dimensions['C'].width = 80.0
    return ws


def write_timeline_header(ws):
    """Created by JXP and Claude

    Write the 15-column (A->O) timeline header plus the
    Calibrations/Red/Blue side block onto a target worksheet, and set
    column widths.

    Inputs
    ------
    ws : openpyxl worksheet to write into.

    Outputs
    -------
    None (ws is modified in place).
    """
    for icol, name in enumerate(TIMELINE_HEADER, start=1):
        cell = ws.cell(row=1, column=icol, value=name)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical='center')

    # Calibrations / Red / Blue side block (cols Q-S), unchecked
    ws['Q2'] = 'Calibrations'
    ws['R2'] = 'Red'
    ws['S2'] = 'Blue'
    for coord in ('Q2', 'R2', 'S2'):
        ws[coord].font = Font(bold=True)
    for irow, label in zip((3, 4, 5), ('Focus', 'Bias', 'Flats')):
        ws.cell(row=irow, column=17, value=label)          # Q
        ws.cell(row=irow, column=18, value=False)          # R (Red)
        ws.cell(row=irow, column=19, value=False)          # S (Blue)

    for letter, width in TIMELINE_WIDTHS.items():
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = 'A2'


def add_target_sheet(wb, targ, first=False):
    """Created by JXP and Claude

    Add one science sheet to wb, named exactly by the target's TNS, with
    the template timeline scaffold: a 'slew to --> ' row and a
    'Science + overhead' row carrying RA/DEC/mags/Comments. Start
    times, Duration, End, red/blue side and Redshift are left blank on
    purpose (exposure times and Kast setup are decided at the telescope).

    Inputs
    ------
    wb : openpyxl.Workbook to receive the sheet.
    targ : pandas Series with TNS, RA_HMS, DEC_DMS, Pri_mag, Sec_mag and
        optionally Pri_POx, FRB_tags.
    first : bool
        If True, add the standard-star bookend note and the
        Standard Star Options reference table.

    Outputs
    -------
    None (wb is modified in place).
    """
    tns = str(targ['TNS'])
    ws = wb.create_sheet(tns)
    write_timeline_header(ws)

    if first:
        note = ws.cell(row=2, column=1, value=STANDARD_NOTE)
        note.font = Font(bold=True, italic=True)

    # Slew row
    ws.cell(row=3, column=6, value='slew to --> ')
    ws.cell(row=3, column=7, value=tns)

    # Science row
    ws.cell(row=4, column=6, value='Science + overhead')
    ws.cell(row=4, column=7, value=tns)
    for col, key in ((8, 'RA_HMS'), (9, 'DEC_DMS')):
        cell = ws.cell(row=4, column=col, value=str(targ[key]))
        cell.number_format = '@'   # keep sexagesimal strings as text
    ws.cell(row=4, column=10, value=float(targ['Pri_mag']))
    ws.cell(row=4, column=11, value=float(targ['Sec_mag']))

    comments = []
    if 'Pri_POx' in targ.index and pd.notna(targ['Pri_POx']):
        comments.append(f"P(O|x)={float(targ['Pri_POx']):.3f}")
    if 'FRB_tags' in targ.index and pd.notna(targ['FRB_tags']):
        comments.append(f"tag={targ['FRB_tags']}")
    if comments:
        ws.cell(row=4, column=15, value='; '.join(comments))

    if first:
        # Standard Star Options reference table (as in the template)
        start = 8
        hdr = ('Standard Star Options', 'RA', 'DEC', 'FK5', 'v')
        for icol, name in enumerate(hdr, start=1):
            cell = ws.cell(row=start, column=icol, value=name)
            cell.font = Font(bold=True)
        for irow, (name, ra, dec, fk5, vmag) in enumerate(
                STANDARD_STARS, start=start + 1):
            ws.cell(row=irow, column=1, value=name)
            for icol, val in ((2, ra), (3, dec)):
                cell = ws.cell(row=irow, column=icol, value=val)
                cell.number_format = '@'
            ws.cell(row=irow, column=4, value=fk5)
            ws.cell(row=irow, column=5, value=vmag)


def add_summary_sheets(wb, targs):
    """Created by JXP and Claude

    Add the 'FRB Observing Summary' and 'Post Observing Tracking'
    sheets, with template headers and the TNS names pre-listed.

    Inputs
    ------
    wb : openpyxl.Workbook to receive the sheets.
    targs : pandas.DataFrame of targets (uses the TNS column, in order).

    Outputs
    -------
    None (wb is modified in place).
    """
    summ = wb.create_sheet('FRB Observing Summary')
    for icol, name in enumerate(
            ('Name', 'Observing Report', 'Target Info', 'z'), start=1):
        cell = summ.cell(row=1, column=icol, value=name)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for irow, tns in enumerate(targs['TNS'], start=2):
        summ.cell(row=irow, column=1, value=str(tns))
    summ.column_dimensions['A'].width = 18.0
    for letter in ('B', 'C'):
        summ.column_dimensions[letter].width = 40.0

    track = wb.create_sheet('Post Observing Tracking')
    hdr = ('TNS', 'Reduction Status', 'Primary', 'Secondary',
           'Uploaded To F4PZ', 'Uploaded To Repo', 'Notes')
    for icol, name in enumerate(hdr, start=1):
        cell = track.cell(row=1, column=icol, value=name)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for irow, tns in enumerate(targs['TNS'], start=2):
        track.cell(row=irow, column=1, value=str(tns))
    track.column_dimensions['A'].width = 18.0
    for letter in ('B', 'C', 'D', 'E', 'F'):
        track.column_dimensions[letter].width = 17.0
    track.column_dimensions['G'].width = 40.0


def make_night_plan(possible_targs_path, template_path, output_path,
                    targets_csv_path=None):
    """Created by JXP and Claude

    Build the full Night-plan workbook and write it to disk.

    Inputs
    ------
    possible_targs_path : str
        Path to the vetted possible_targs xlsx.
    template_path : str
        Path to the template night-plan xlsx (read-only; supplies the
        Observing Checklist).
    output_path : str
        Path of the xlsx workbook to write.
    targets_csv_path : str or None
        Optional targets CSV for P(O|x) / FRB tags in Comments.

    Outputs
    -------
    output_path (str), after writing the workbook. Sheet order:
    Observing Checklist, one sheet per target (named by TNS, RA order),
    FRB Observing Summary, Post Observing Tracking.
    """
    targs = read_targets(possible_targs_path, targets_csv_path)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # drop the default empty sheet

    copy_checklist(template_path, wb)
    for idx, (_, targ) in enumerate(targs.iterrows()):
        add_target_sheet(wb, targ, first=(idx == 0))
    add_summary_sheets(wb, targs)

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    wb.save(output_path)
    return output_path


def main():
    """Created by JXP and Claude

    Command-line driver. Inputs are the argparse options below; the
    output is the Night-plan xlsx written to --output.
    """
    base = '/home/xavier/Projects/night-planner/Night_plans/Lick-2026B-01'
    parser = argparse.ArgumentParser(
        description='Build a Lick Shane/Kast Night-plan workbook with one '
                    'sheet per target (Created by JXP and Claude).')
    parser.add_argument(
        '--possible_targs',
        default=os.path.join(base, 'Lick-2026B-01_possible_targs.xlsx'),
        help='Vetted possible_targs xlsx (sheet possible_targs).')
    parser.add_argument(
        '--targets_csv',
        default=os.path.join(base, 'Lick-2026B-01_targets.csv'),
        help='Optional targets CSV with Pri_POx / FRB_tags.')
    parser.add_argument(
        '--template',
        default=('/mnt/scratch/xavier/Observing/Observations/Lick_Kast/'
                 'obs_docs/2026/Lick-2026A-4/Lick_2026A-4_Night_plan.xlsx'),
        help='Template night-plan xlsx (read-only).')
    parser.add_argument(
        '--output',
        default=os.path.join(base, 'Lick_2026B-01_Night_plan.xlsx'),
        help='Output Night-plan xlsx path.')
    args = parser.parse_args()

    out = make_night_plan(args.possible_targs, args.template, args.output,
                          targets_csv_path=args.targets_csv)
    print(f'Wrote {out}')


if __name__ == '__main__':
    main()
