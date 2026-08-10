# Created by JXP and Claude
"""Created by JXP and Claude

Estimate Kast exposure times for the targets of a Lick Night plan and
write them into the Night-plan workbook (in place).

Kast setup assumed (2026B-01): d57 dichroic, red 600/7500 grating,
blue 600/4310 grism, 2" slit.  Kast is a dual-beam spectrograph -- the
red and blue arms expose SIMULTANEOUSLY through the dichroic -- so the
per-target wall-clock science time is ONE arm's total integration (the
arms are balanced so red_total ~= blue_total), not their sum.

The exposure heuristic is empirical: it was calibrated against the
authors' 9 previous Lick/Kast night-plan workbooks (2024-2026A), in
which total on-source integration scales with the primary-host
r-magnitude, the red arm is split into ~600 s subframes, and the less
efficient blue arm into longer (900-1800 s) subframes.

Inputs
------
- possible_targs xlsx (sheet 'possible_targs'): columns
  TNS, RA_HMS, DEC_DMS, Epoch, Pri_mag, Sec_mag.
- Night-plan xlsx (as built by make_lick_night_plan.py): one timeline
  sheet per target, named by TNS, with a 'Science + overhead' row and
  columns 'Duration (hh:min)' (D), 'red side' (L), 'blue side' (M).

Output
------
- The Night-plan xlsx updated IN PLACE: per target sheet, the
  'Science + overhead' row gets red side, blue side and Duration; the
  'slew to --> ' row Duration is kept at 5 min; and the Kast setup is
  stamped in cell Q1.  With --dry_run, only a table is printed.
"""

import argparse
import datetime

import pandas as pd
import openpyxl
from openpyxl.styles import Font

# Kast setup for this run, stamped on each target sheet
KAST_SETUP_NOTE = ('Kast setup: d57 dichroic | red 600/7500 | '
                   'blue 600/4310 | 2" slit')

# Timeline column letters (fixed by the template; see
# make_lick_night_plan.py TIMELINE_HEADER)
COL_DURATION = 'D'
COL_ACTION = 'F'
COL_RED = 'L'
COL_BLUE = 'M'

SLEW_SECONDS = 300  # keep the 'slew to --> ' row at 5 min


def estimate_exposure(mag):
    """Created by JXP and Claude

    Estimate the Kast red/blue exposure sequence for a target from its
    primary-host r-band magnitude.

    Empirical basis: deterministic mapping calibrated to the authors'
    9 prior Lick/Kast night plans (e.g. a mag ~18.6 host in 2026A-2
    was observed 6 x 600 red / 3 x 1200 blue).  Total integration T per
    arm scales with magnitude:
        mag < 15.5          -> T = 1800 s
        15.5 <= mag < 17.5  -> T = 2400 s
        17.5 <= mag < 19.3  -> T = 3600 s
        mag >= 19.3         -> T = 5400 s
    The red arm is split into 600 s subframes (300 s for very bright
    hosts, mag < 13); the less efficient blue arm into longer subframes
    (900 s for T <= 1800; 1200 s for 2400 <= T <= 3600; 1800 s for
    T >= 5400), keeping red_total ~= blue_total.  At least 2 subframes
    per arm (cosmic-ray rejection).  Because Kast is dual-beam, the
    wall clock is the longer of the two arm totals, not their sum.

    Inputs
    ------
    mag : float
        Primary-host r-band magnitude.

    Returns
    -------
    red_str : str
        Red-side sequence, e.g. '6 x 600'.
    blue_str : str
        Blue-side sequence, e.g. '3 x 1200'.
    duration_seconds : int
        Wall-clock science duration = max(red_total, blue_total).
    """
    # Total integration per arm from the host magnitude
    if mag < 15.5:
        T = 1800
    elif mag < 17.5:
        T = 2400
    elif mag < 19.3:
        T = 3600
    else:
        T = 5400

    # Red subframes: 600 s standard, 300 s for very bright hosts
    red_sub = 300 if mag < 13 else 600
    n_red = max(2, round(T / red_sub))

    # Blue subframes: longer, since the blue arm is less efficient
    if T <= 1800:
        blue_sub = 900
    elif T <= 3600:
        blue_sub = 1200
    else:
        blue_sub = 1800
    n_blue = max(2, round(T / blue_sub))

    red_str = f'{n_red} x {red_sub}'
    blue_str = f'{n_blue} x {blue_sub}'
    # Dual-beam: arms expose simultaneously -> wall clock is the max
    duration_seconds = max(n_red * red_sub, n_blue * blue_sub)

    return red_str, blue_str, duration_seconds


def seconds_to_time(seconds):
    """Created by JXP and Claude

    Convert a duration in seconds to a datetime.time for an Excel
    hh:mm cell.

    Inputs
    ------
    seconds : int
        Duration in seconds (< 24 hours).

    Returns
    -------
    datetime.time
        The duration expressed as a time of day (h, m).
    """
    minutes = int(round(seconds / 60.))
    return datetime.time(minutes // 60, minutes % 60)


def find_action_row(ws, action):
    """Created by JXP and Claude

    Find the timeline row whose Action column matches a given label.

    Inputs
    ------
    ws : openpyxl worksheet
        A target timeline sheet.
    action : str
        Action label to match (stripped, e.g. 'Science + overhead').

    Returns
    -------
    int or None
        The 1-indexed row number, or None if not found.
    """
    for row in range(1, ws.max_row + 1):
        val = ws[f'{COL_ACTION}{row}'].value
        if isinstance(val, str) and val.strip() == action:
            return row
    return None


def update_night_plan(possible_targs_path, night_plan_path, dry_run=False):
    """Created by JXP and Claude

    Fill Kast exposure estimates into a Night-plan workbook, in place.

    For each target in the possible_targs sheet, locate its TNS-named
    sheet in the Night plan and, on the 'Science + overhead' row, set
    the red side, blue side and Duration (hh:min) cells from
    estimate_exposure(Pri_mag).  'Redshift;' is left blank (unknown
    pre-observation).  The 'slew to --> ' row Duration is set to 5 min,
    and the Kast setup note is stamped in cell Q1 of each target sheet.

    Inputs
    ------
    possible_targs_path : str
        Path to the possible_targs xlsx (sheet 'possible_targs').
    night_plan_path : str
        Path to the Night-plan xlsx, updated IN PLACE (openpyxl
        load -> edit -> save; all other content is preserved).
    dry_run : bool
        If True, print the exposure table and do not write the xlsx.

    Returns
    -------
    pandas.DataFrame
        Table of TNS, Pri_mag, red side, blue side, Duration.
    """
    targs = pd.read_excel(possible_targs_path, sheet_name='possible_targs')

    rows = []
    for _, targ in targs.iterrows():
        red_str, blue_str, dur_s = estimate_exposure(targ.Pri_mag)
        rows.append(dict(TNS=targ.TNS, Pri_mag=targ.Pri_mag,
                         red_side=red_str, blue_side=blue_str,
                         duration=f'{dur_s // 3600}:{(dur_s % 3600) // 60:02d}'))
    table = pd.DataFrame(rows)
    print(table.to_string(index=False))

    if dry_run:
        print('\n--dry_run: Night plan NOT modified')
        return table

    wb = openpyxl.load_workbook(night_plan_path)
    for row in rows:
        tns = row['TNS']
        if tns not in wb.sheetnames:
            print(f'WARNING: no sheet for {tns}; skipping')
            continue
        ws = wb[tns]

        # Science + overhead row: Duration, red side, blue side
        sci_row = find_action_row(ws, 'Science + overhead')
        if sci_row is None:
            print(f'WARNING: no Science + overhead row on {tns}; skipping')
            continue
        red_str, blue_str, dur_s = estimate_exposure(
            targs.loc[targs.TNS == tns, 'Pri_mag'].values[0])
        ws[f'{COL_RED}{sci_row}'] = red_str
        ws[f'{COL_BLUE}{sci_row}'] = blue_str
        dur_cell = ws[f'{COL_DURATION}{sci_row}']
        dur_cell.value = seconds_to_time(dur_s)
        dur_cell.number_format = 'h:mm'
        # 'Redshift;' column deliberately left blank (unknown pre-obs)

        # Keep the slew row at 5 min
        slew_row = find_action_row(ws, 'slew to -->')
        if slew_row is not None:
            slew_cell = ws[f'{COL_DURATION}{slew_row}']
            slew_cell.value = seconds_to_time(SLEW_SECONDS)
            slew_cell.number_format = 'h:mm'

        # Stamp the Kast setup once per sheet, top-right of the timeline
        setup_cell = ws['Q1']
        setup_cell.value = KAST_SETUP_NOTE
        setup_cell.font = Font(bold=True)

    wb.save(night_plan_path)
    print(f'\nUpdated in place: {night_plan_path}')
    return table


def main():
    """Created by JXP and Claude

    Command-line driver: parse arguments and update the Night plan.

    Inputs (argparse)
    -----------------
    --possible_targs : str
        Path to the possible_targs xlsx.
    --night_plan : str
        Path to the Night-plan xlsx to update in place.
    --dry_run : flag
        Print the exposure table without writing the workbook.

    Returns
    -------
    None
    """
    parser = argparse.ArgumentParser(
        description='Estimate Kast exposure times (d57, 600/7500 red, '
                    '600/4310 blue, 2" slit) and fill the Night plan')
    parser.add_argument(
        '--possible_targs', type=str,
        default='Night_plans/Lick-2026B-01/Lick-2026B-01_possible_targs.xlsx',
        help='possible_targs xlsx (sheet possible_targs)')
    parser.add_argument(
        '--night_plan', type=str,
        default='Night_plans/Lick-2026B-01/Lick_2026B-01_Night_plan.xlsx',
        help='Night-plan xlsx to update in place')
    parser.add_argument(
        '--dry_run', action='store_true',
        help='print the exposure table without writing the workbook')
    args = parser.parse_args()

    update_night_plan(args.possible_targs, args.night_plan,
                      dry_run=args.dry_run)


if __name__ == '__main__':
    main()
