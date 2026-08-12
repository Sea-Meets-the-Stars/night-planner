"""
Build night plan Excel file for Lick 2026B-01

Created by JXP and Claude
"""

import pandas as pd
import numpy as np
from datetime import datetime, timedelta, timezone
import pytz
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from astropy import units as u
from astropy.coordinates import SkyCoord, EarthLocation
from astropy.time import Time

# Constants
LICK = EarthLocation(lat=37.3414*u.deg, lon=-121.6429*u.deg, height=1283*u.m)
# Lick calendar uses PST (UTC-8) year-round
PST = timezone(timedelta(hours=-8))

def pst_to_lst(pst_dt):
    """Convert PST (UTC-8) datetime to LST at Lick."""
    if pst_dt.tzinfo is None:
        pst_dt = pst_dt.replace(tzinfo=PST)
    utc_dt = pst_dt.astimezone(timezone.utc)
    t = Time(utc_dt)
    lst = t.sidereal_time('apparent', longitude=LICK.lon)
    return lst.hour

def pst_to_ut(pst_dt):
    """Convert PST (UTC-8) datetime to UT (decimal hours)."""
    if pst_dt.tzinfo is None:
        pst_dt = pst_dt.replace(tzinfo=PST)
    utc_dt = pst_dt.astimezone(timezone.utc)
    return utc_dt.hour + utc_dt.minute/60.0 + utc_dt.second/3600.0

def format_ra_dec(ra_deg, dec_deg):
    """Format RA/Dec in sexagesimal."""
    coord = SkyCoord(ra=ra_deg*u.deg, dec=dec_deg*u.deg, frame='icrs')
    ra_str = coord.ra.to_string(unit=u.hour, sep=':', precision=2, pad=True)
    dec_str = coord.dec.to_string(unit=u.deg, sep=':', precision=1, pad=True)
    return ra_str, dec_str

def add_observing_checklist(ws):
    """Add observing checklist sheet."""
    checklist_items = [
        "Dome open",
        "Telescope tracking",
        "Focus achieved",
        "Standard star (evening)",
        "Science targets started",
        "Standard star (morning)",
        "Bias frames taken",
        "Flat fields taken",
        "Dome closed",
        "Data backed up"
    ]

    ws['A1'] = "Lick 2026B-01 Observing Checklist"
    ws['A1'].font = Font(size=14, bold=True)

    ws['A3'] = "Task"
    ws['B3'] = "Done"
    ws['A3'].font = Font(bold=True)
    ws['B3'].font = Font(bold=True)

    for i, item in enumerate(checklist_items, start=4):
        ws[f'A{i}'] = item
        ws[f'B{i}'] = False

    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 10

def add_night_sheet(wb, date_str, selected_targets, twilight, start_pst):
    """
    Add the main observing night sheet.

    Parameters
    ----------
    wb : Workbook
    date_str : str (e.g., "August 13th")
    selected_targets : DataFrame
    twilight : dict
    start_pst : datetime
    """
    ws = wb.create_sheet(date_str)

    # Set up column headers
    headers = ['Start (LST)', 'Start (UT)', 'Start (PST)', 'Duration', 'End (PST)',
               'Action', 'target', 'RA', 'DEC', 'HA start', 'HA end', 'mag(r)', 'sec mag',
               'red side', 'blue side', 'Redshift;', 'Comments']

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col)
        cell.value = header
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")

    # Set column widths
    ws.column_dimensions['A'].width = 12  # LST
    ws.column_dimensions['B'].width = 12  # UT
    ws.column_dimensions['C'].width = 12  # PST
    ws.column_dimensions['D'].width = 12  # Duration
    ws.column_dimensions['E'].width = 12  # End
    ws.column_dimensions['F'].width = 20  # Action
    ws.column_dimensions['G'].width = 15  # target
    ws.column_dimensions['H'].width = 12  # RA
    ws.column_dimensions['I'].width = 12  # DEC
    ws.column_dimensions['J'].width = 10  # HA start
    ws.column_dimensions['K'].width = 10  # HA end
    ws.column_dimensions['L'].width = 8   # mag(r)
    ws.column_dimensions['M'].width = 8   # sec mag
    ws.column_dimensions['N'].width = 15  # red side
    ws.column_dimensions['O'].width = 15  # blue side
    ws.column_dimensions['P'].width = 10  # Redshift
    ws.column_dimensions['Q'].width = 30  # Comments

    current_row = 2
    current_time = start_pst

    # Function to add a row
    def add_row(action, target='', ra='', dec='', mag_r='', sec_mag='',
                red_side='', blue_side='', redshift='', comments='',
                duration_min=0, ra_hours=None):
        nonlocal current_row, current_time

        # Calculate times
        lst_start = pst_to_lst(current_time)
        ut = pst_to_ut(current_time)
        pst_str = current_time.strftime('%H:%M')
        end_time = current_time + timedelta(minutes=duration_min)
        end_str = end_time.strftime('%H:%M')
        lst_end = pst_to_lst(end_time) if duration_min > 0 else lst_start

        # Calculate hour angles if RA provided
        ha_start_str = ''
        ha_end_str = ''
        if ra_hours is not None:
            # HA = LST - RA, wrapped to ±12h
            ha_start = lst_start - ra_hours
            if ha_start > 12:
                ha_start -= 24
            elif ha_start < -12:
                ha_start += 24

            ha_end = lst_end - ra_hours
            if ha_end > 12:
                ha_end -= 24
            elif ha_end < -12:
                ha_end += 24

            ha_start_str = f"{ha_start:.2f}"
            ha_end_str = f"{ha_end:.2f}"

        # Format duration
        if duration_min > 0:
            duration_str = f"{duration_min} min"
        else:
            duration_str = ""

        # Write row
        ws.cell(row=current_row, column=1, value=f"{lst_start:.2f}")
        ws.cell(row=current_row, column=2, value=f"{ut:.2f}")
        ws.cell(row=current_row, column=3, value=pst_str)
        ws.cell(row=current_row, column=4, value=duration_str)
        ws.cell(row=current_row, column=5, value=end_str if duration_min > 0 else '')
        ws.cell(row=current_row, column=6, value=action)
        ws.cell(row=current_row, column=7, value=target)
        ws.cell(row=current_row, column=8, value=ra)
        ws.cell(row=current_row, column=9, value=dec)
        ws.cell(row=current_row, column=10, value=ha_start_str)
        ws.cell(row=current_row, column=11, value=ha_end_str)
        ws.cell(row=current_row, column=12, value=mag_r)
        ws.cell(row=current_row, column=13, value=sec_mag)
        ws.cell(row=current_row, column=14, value=red_side)
        ws.cell(row=current_row, column=15, value=blue_side)
        ws.cell(row=current_row, column=16, value=redshift)
        ws.cell(row=current_row, column=17, value=comments)

        current_row += 1
        current_time = end_time

    # Evening standard star
    # HZ_44: RA 13:23:35, Dec +36:07:59, g=11.7
    add_row(action="Standard Star", target="HZ_44",
            ra="13:23:35", dec="+36:07:59",
            mag_r="11.7", sec_mag="",
            red_side="1 x 60", blue_side="1 x 60",
            comments="Evening standard",
            duration_min=5,
            ra_hours=13.393)  # 13:23:35 in hours

    # Focus
    add_row(action="Focus", duration_min=5)

    # Science targets
    for idx, targ in selected_targets.iterrows():
        tns = targ['TNS']
        ra_deg = targ['RA_deg']
        dec_deg = targ['Dec_deg']
        ra_hours = targ['RA_hours']
        mag_r = targ['mag_r']
        pox = targ['POx']
        tags = targ['tags']

        ra_str, dec_str = format_ra_dec(ra_deg, dec_deg)

        # Determine exposure times
        if mag_r <= 15:
            red_exp = "2 x 900"
            blue_exp = "1 x 1800"
            total_min = 40
        elif mag_r <= 17:
            red_exp = "3 x 900"
            blue_exp = "2 x 1350"
            total_min = 60
        elif mag_r <= 19:
            red_exp = "4 x 900"
            blue_exp = "2 x 1800"
            total_min = 80
        else:
            red_exp = "6 x 900"
            blue_exp = "3 x 1800"
            total_min = 120

        # Slew
        add_row(action=f"slew to --> {tns}", duration_min=2)

        # Science observation
        add_row(action="Science + overhead",
                target=tns,
                ra=ra_str,
                dec=dec_str,
                mag_r=f"{mag_r:.2f}",
                sec_mag="",
                red_side=red_exp,
                blue_side=blue_exp,
                redshift="",
                comments=f"P(O|x)={pox:.3f}; tags={tags}",
                duration_min=total_min,
                ra_hours=ra_hours)

    # Morning standard star
    add_row(action="Standard Star", target="Feige_110",
            ra="23:19:58.4", dec="-05:09:56",
            mag_r="11.8", sec_mag="",
            red_side="1 x 60", blue_side="1 x 60",
            comments="Morning standard",
            duration_min=5,
            ra_hours=23.333)  # 23:19:58.4 in hours

    # Calibrations
    add_row(action="Bias frames", duration_min=5)
    add_row(action="Flat fields", duration_min=10)

    # Add note about standard star options
    note_row = current_row + 2
    ws.cell(row=note_row, column=1, value="Standard Star Options:")
    ws.cell(row=note_row, column=1).font = Font(bold=True)

    std_stars = [
        ("HZ_44", "13:23:35", "+36:07:59", "11.7"),
        ("Feige_110", "23:19:58.4", "-05:09:56", "11.8"),
        ("BD+28_4211", "21:51:11", "+28:51:50", "10.5")
    ]

    for i, (name, ra, dec, mag) in enumerate(std_stars, start=note_row+1):
        ws.cell(row=i, column=1, value=name)
        ws.cell(row=i, column=2, value=ra)
        ws.cell(row=i, column=3, value=dec)
        ws.cell(row=i, column=4, value=mag)

    # Add sidereal time calculator section (columns S-T, after HA columns J-K)
    calc_row = 1
    ws.cell(row=calc_row, column=19, value="Sidereal Time Calculator")
    ws.cell(row=calc_row, column=19).font = Font(bold=True)

    calc_row += 1
    ws.cell(row=calc_row, column=19, value="PST")
    ws.cell(row=calc_row, column=20, value="LST")
    ws.cell(row=calc_row, column=19).font = Font(bold=True)
    ws.cell(row=calc_row, column=20).font = Font(bold=True)

    # Add a few sample times
    sample_times = [
        start_pst,
        start_pst + timedelta(hours=2),
        start_pst + timedelta(hours=4),
        start_pst + timedelta(hours=6),
    ]

    for i, sample_time in enumerate(sample_times, start=calc_row+1):
        pst_str = sample_time.strftime('%H:%M')
        lst = pst_to_lst(sample_time)
        ws.cell(row=i, column=19, value=pst_str)
        ws.cell(row=i, column=20, value=f"{lst:.2f}")

    return ws

def build_night_plan(output_path, selected_csv, date='2026-08-13'):
    """
    Build the complete night plan Excel file.

    Parameters
    ----------
    output_path : str
        Path to output Excel file
    selected_csv : str
        Path to selected targets CSV
    date : str
        Observation date
    """
    # Load selected targets
    selected_targets = pd.read_csv(selected_csv)

    # Parse twilight times from the min_am_time_pst column
    # (This is a hack - ideally we'd recalculate or pass twilight explicitly)
    sample_time_str = selected_targets.iloc[0]['min_am_time_pst']
    sample_time = pd.to_datetime(sample_time_str)

    # For Aug 13, 2026, from our calculation:
    # Evening 18° twilight: 20:38 PST (UTC-8)
    # Timeline start: 20:30 PST (rounded down to half-hour)
    start_pst = datetime(2026, 8, 13, 20, 30, 0, tzinfo=PST)

    # Create workbook
    wb = Workbook()
    wb.remove(wb.active)  # Remove default sheet

    # Add sheets
    checklist = wb.create_sheet("Observing Checklist")
    add_observing_checklist(checklist)

    twilight = {}  # Placeholder
    add_night_sheet(wb, "August 13th", selected_targets, twilight, start_pst)

    # Save
    wb.save(output_path)
    print(f"Night plan saved to: {output_path}")

def main():
    """Main execution."""
    selected_csv = '/home/lordrick/Projects/night-planner/Night_plans/Lick-2026B-01/selected_targets_night1.csv'
    output_path = '/home/lordrick/Projects/night-planner/Night_plans/Lick-2026B-01/Lick_2026B-01_Night_plan.xlsx'

    build_night_plan(output_path, selected_csv)

if __name__ == '__main__':
    main()
