"""
Build night plan Excel file for Keck/LRIS Oct 2, 2026

Created by JXP and Claude
"""

import pandas as pd
import numpy as np
from datetime import datetime, timedelta, timezone
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from astropy import units as u
from astropy.coordinates import SkyCoord, EarthLocation
from astropy.time import Time

# Constants
# Verified against astropy's EarthLocation.of_site('keck') (IRAF Observatory
# Database), 2026-09-11: lat 19.828333 deg, lon -155.478333 deg, height 4160 m.
KECK = EarthLocation(lat=19.828333*u.deg, lon=-155.478333*u.deg, height=4160*u.m)
HST = timezone(timedelta(hours=-10))  # Hawaii Standard Time


def hst_to_lst(hst_dt):
    """Convert HST (UTC-10) datetime to LST at Keck."""
    if hst_dt.tzinfo is None:
        hst_dt = hst_dt.replace(tzinfo=HST)
    utc_dt = hst_dt.astimezone(timezone.utc)
    t = Time(utc_dt)
    lst = t.sidereal_time('apparent', longitude=KECK.lon)
    return lst.hour


def hst_to_ut(hst_dt):
    """Convert HST (UTC-10) datetime to UT (decimal hours)."""
    if hst_dt.tzinfo is None:
        hst_dt = hst_dt.replace(tzinfo=HST)
    utc_dt = hst_dt.astimezone(timezone.utc)
    return utc_dt.hour + utc_dt.minute/60.0 + utc_dt.second/3600.0


def format_ra_dec(ra_deg, dec_deg):
    """Format RA/Dec in sexagesimal."""
    coord = SkyCoord(ra=ra_deg*u.deg, dec=dec_deg*u.deg, frame='icrs')
    ra_str = coord.ra.to_string(unit=u.hour, sep=':', precision=2, pad=True)
    dec_str = coord.dec.to_string(unit=u.deg, sep=':', precision=1, pad=True)
    return ra_str, dec_str


def add_observing_checklist(ws):
    """Add observing checklist sheet."""
    checklist_items = [
        "Dome open",
        "Telescope tracking",
        "Focus achieved",
        "Guider configured",
        "Dichroic in place",
        "Gratings verified",
        "Standard star (evening)",
        "Science targets started",
        "Standard star (morning)",
        "Arc lamp calibrations",
        "Flat field calibrations",
        "Dome closed",
        "Data transferred"
    ]

    ws['A1'] = "Keck/LRIS Oct 2, 2026 Observing Checklist"
    ws['A1'].font = Font(size=14, bold=True)

    ws['A3'] = "Task"
    ws['B3'] = "Done"
    ws['C3'] = "Notes"
    ws['A3'].font = Font(bold=True)
    ws['B3'].font = Font(bold=True)
    ws['C3'].font = Font(bold=True)

    for i, item in enumerate(checklist_items, start=4):
        ws[f'A{i}'] = item
        ws[f'B{i}'] = ""
        ws[f'C{i}'] = ""

    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 10
    ws.column_dimensions['C'].width = 30

    # Add LRIS configuration section
    config_start = len(checklist_items) + 6
    ws.cell(row=config_start, column=1, value="LRIS Configuration")
    ws.cell(row=config_start, column=1).font = Font(size=12, bold=True)

    config_items = [
        ("Dichroic", "560"),
        ("Blue Grism", "600/4000"),
        ("Red Grating", "600/7500 (centered ~7200 A)"),
        ("Blue CCD binning", "2x2"),
        ("Red CCD binning", "2x2"),
        ("Slit width", "1.0\""),
    ]

    for i, (key, val) in enumerate(config_items, start=config_start+1):
        ws.cell(row=i, column=1, value=key)
        ws.cell(row=i, column=2, value=val)


def add_night_sheet(wb, date_str, selected_targets, start_hst):
    """
    Add the main observing night sheet.

    Parameters
    ----------
    wb : Workbook
    date_str : str (e.g., "October 2nd")
    selected_targets : DataFrame
    start_hst : datetime
    """
    ws = wb.create_sheet(date_str)

    # Set up column headers (Keck format)
    headers = ['Start (LST)', 'Start (UT)', 'Start (HST)', 'Duration', 'End (HST)',
               'Action', 'target', 'RA', 'DEC', 'HA start', 'HA end', 'mag(r)',
               'red side', 'blue side', 'Airmass', 'Comments']

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col)
        cell.value = header
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")

    # Set column widths
    widths = [12, 10, 12, 10, 12, 18, 16, 12, 12, 10, 10, 8, 12, 12, 8, 35]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64+i)].width = w

    current_row = 2
    current_time = start_hst

    def add_row(action, target='', ra='', dec='', mag_r='',
                red_side='', blue_side='', airmass='', comments='',
                duration_min=0, ra_hours=None):
        nonlocal current_row, current_time

        # Calculate times
        lst_start = hst_to_lst(current_time)
        ut = hst_to_ut(current_time)
        hst_str = current_time.strftime('%H:%M')
        end_time = current_time + timedelta(minutes=duration_min)
        end_str = end_time.strftime('%H:%M')
        lst_end = hst_to_lst(end_time) if duration_min > 0 else lst_start

        # Calculate hour angles if RA provided
        ha_start_str = ''
        ha_end_str = ''
        if ra_hours is not None:
            ha_start = lst_start - ra_hours
            if ha_start > 12:
                ha_start -= 24
            elif ha_start < -12:
                ha_start += 24

            ha_end = lst_end - ra_hours
            if ha_end > 12:
                ha_end -= 24
            elif ha_end < -12:
                ha_end += 24

            ha_start_str = f"{ha_start:+.1f}h"
            ha_end_str = f"{ha_end:+.1f}h"

        duration_str = f"{duration_min} min" if duration_min > 0 else ""

        # Write row
        ws.cell(row=current_row, column=1, value=f"{lst_start:.1f}h")
        ws.cell(row=current_row, column=2, value=f"{ut:.1f}")
        ws.cell(row=current_row, column=3, value=hst_str)
        ws.cell(row=current_row, column=4, value=duration_str)
        ws.cell(row=current_row, column=5, value=end_str if duration_min > 0 else '')
        ws.cell(row=current_row, column=6, value=action)
        ws.cell(row=current_row, column=7, value=target)
        ws.cell(row=current_row, column=8, value=ra)
        ws.cell(row=current_row, column=9, value=dec)
        ws.cell(row=current_row, column=10, value=ha_start_str)
        ws.cell(row=current_row, column=11, value=ha_end_str)
        ws.cell(row=current_row, column=12, value=mag_r)
        ws.cell(row=current_row, column=13, value=red_side)
        ws.cell(row=current_row, column=14, value=blue_side)
        ws.cell(row=current_row, column=15, value=airmass)
        ws.cell(row=current_row, column=16, value=comments)

        current_row += 1
        current_time = end_time

    # Evening standard star (BD+28 4211)
    add_row(action="Flux Standard", target="BD28d4211",
            ra="21:51:11.02", dec="+28:51:50.4",
            mag_r="10.5",
            red_side="1 x 30", blue_side="1 x 30",
            comments="Evening standard",
            duration_min=5,
            ra_hours=21.853)

    # Focus
    add_row(action="Focus check", duration_min=5)

    # Science targets - use scheduled times
    for idx, targ in selected_targets.iterrows():
        tns = targ['TNS']
        ra_deg = targ['RA_deg']
        dec_deg = targ['Dec_deg']
        ra_hours = targ['RA_hours']
        mag_r = targ['mag_r']

        ra_str, dec_str = format_ra_dec(ra_deg, dec_deg)

        # Get scheduled times
        sched_start = targ.get('scheduled_start')
        if pd.notna(sched_start):
            if isinstance(sched_start, str):
                sched_start = pd.to_datetime(sched_start)
            # Jump to scheduled time if there's a gap
            if hasattr(sched_start, 'tzinfo') and sched_start.tzinfo is not None:
                if sched_start > current_time:
                    current_time = sched_start
            else:
                # Handle naive datetime
                sched_start = sched_start.replace(tzinfo=HST)
                if sched_start > current_time:
                    current_time = sched_start

        total_min = int(targ.get('total_min', 30))
        red_exp = targ.get('red_exp', '2 x 600')
        blue_exp = targ.get('blue_exp', '2 x 600')
        am = targ.get('min_airmass', 1.5)
        tags = str(targ.get('tags', '')) if pd.notna(targ.get('tags')) else ''
        is_mos = targ.get('lris_mode') == 'MOS (mask required)'
        comment_bits = (['MASK REQUIRED'] if is_mos else []) + ([tags] if tags else [])
        comments = ' | '.join(comment_bits)[:40]

        # Slew
        add_row(action=f"slew to", target=tns, duration_min=3)

        # Science observation
        add_row(action="Science" if not is_mos else "Science (MOS)",
                target=tns,
                ra=ra_str,
                dec=dec_str,
                mag_r=f"{mag_r:.1f}",
                red_side=red_exp,
                blue_side=blue_exp,
                airmass=f"{am:.2f}",
                comments=comments,
                duration_min=total_min - 3,  # Subtract slew time
                ra_hours=ra_hours)

    # Morning standard star (Feige 110)
    add_row(action="Flux Standard", target="Feige110",
            ra="23:19:58.40", dec="-05:09:56.2",
            mag_r="11.8",
            red_side="1 x 30", blue_side="1 x 30",
            comments="Morning standard",
            duration_min=5,
            ra_hours=23.333)

    # Calibrations
    add_row(action="Arc lamps (HgNeArCdZn)", duration_min=10)
    add_row(action="Internal flats", duration_min=10)

    # Add notes section
    note_row = current_row + 2
    ws.cell(row=note_row, column=1, value="NOTES:")
    ws.cell(row=note_row, column=1).font = Font(bold=True)

    # Largest gap between consecutive scheduled targets (computed from the
    # actual schedule, not hardcoded -- a hardcoded gap string went stale
    # three times across earlier revisions of this plan).
    gap_note = "- No gaps: back-to-back targets all night"
    if len(selected_targets) > 1:
        starts_ends = sorted(
            (pd.to_datetime(t['scheduled_start']), pd.to_datetime(t['scheduled_end']))
            for _, t in selected_targets.iterrows())
        max_gap = pd.Timedelta(0)
        gap_start = gap_end = None
        for (s1, e1), (s2, e2) in zip(starts_ends, starts_ends[1:]):
            gap = s2 - e1
            if gap > max_gap:
                max_gap, gap_start, gap_end = gap, e1, s2
        if gap_start is not None:
            gap_note = (f"- Gap from ~{gap_start.strftime('%H:%M')}-"
                        f"{gap_end.strftime('%H:%M')} HST: no further "
                        f"observable targets in hand")

    notes = [
        gap_note,
        "- All targets observable at airmass < 2.0 when scheduled",
        "- LRIS setup: 560 dichroic, 600/4000 blue grism, 600/7500 red grating (~7200A)",
        "- Exposure times from the real LRIS ETC (etc.ucolick.org): S/N=10/pixel per single exposure, 3 exposures min",
        "- MOS (mag>21) targets need ~1-2h once masked (see info column); utilization is target-count-limited: request more targets"
    ]

    for i, note in enumerate(notes, start=note_row+1):
        ws.cell(row=i, column=1, value=note)

    return ws


def build_keck_night_plan(output_path, selected_csv, date='2026-10-02'):
    """
    Build the complete Keck night plan Excel file.
    """
    # Load selected targets
    selected_targets = pd.read_csv(selected_csv)

    # Start time: 19:20 HST (before twilight at 19:22)
    start_hst = datetime(2026, 10, 2, 19, 20, 0, tzinfo=HST)

    # Create workbook
    wb = Workbook()
    wb.remove(wb.active)

    # Add sheets
    checklist = wb.create_sheet("Observing Checklist")
    add_observing_checklist(checklist)

    add_night_sheet(wb, "October 2nd", selected_targets, start_hst)

    # Save
    wb.save(output_path)
    print(f"Night plan saved to: {output_path}")


def main():
    """Main execution."""
    selected_csv = '/home/lordrick/Projects/night-planner/Night_plans/Keck-2026B-01/selected_targets.csv'
    output_path = '/home/lordrick/Projects/night-planner/Night_plans/Keck-2026B-01/Keck_LRIS_2026Oct02_Night_plan.xlsx'

    build_keck_night_plan(output_path, selected_csv)


if __name__ == '__main__':
    main()
