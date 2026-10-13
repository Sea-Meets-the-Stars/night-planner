"""
Add the Night-2 (2026-09-15) sheets to the Lick-2026B-02 Night-plan
workbook, in place.

Created by JXP and Claude

Adds two sheets to the existing Lick_2026B-02_Night_plan.xlsx (alongside
"Observing Checklist", "September 14th", "Extras (Gap-Fill)" from night 1):
  - "September 15th": the FRB priority timeline -- repeats of everything
    observed night 1 that did NOT yield a redshift (FRB20260623C and
    FRB20260507D got z's and are excluded; see
    scripts/lick_2026b02_night2_calculations.py).
  - "Spare Targets (Sep 15)": the author-supplied non-FRB filler targets
    scheduled into the FRB schedule's idle gaps, plus a note on the two
    that are never Shane-HA-legal tonight (2026yly, 2026sqf) with the
    real numbers on when their window would open relative to sunrise.
"""

import sys
from datetime import timedelta, timezone

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from astropy import units as u
from astropy.coordinates import SkyCoord

sys.path.insert(0, '/home/lordrick/Projects/night-planner')
from scripts.lick_2026b02_calculations import find_twilight, calculate_lst, ha_at
from scripts.lick_2026b02_night2_calculations import (
    DATE, OUT_DIR, never_legal_diagnostic, REPEAT_PRIMARY_TNS, REPEAT_EXTRA_TNS,
)

WB_PATH = f'{OUT_DIR}/Lick_2026B-02_Night_plan.xlsx'
SHEET_NAME = 'September 15th'
SPARE_SHEET_NAME = 'Spare Targets (Sep 15)'

STD_STAR = dict(name='Feige_110', ra='23:19:58.4', dec='-05:09:56',
                 ra_hours=23.333, mag='11.8')


def pst_to_lst(pst_dt):
    return calculate_lst(pst_dt.astimezone(timezone.utc)).hour


def pst_to_ut(pst_dt):
    u_ = pst_dt.astimezone(timezone.utc)
    return u_.hour + u_.minute / 60.0 + u_.second / 3600.0


def format_ra_dec(ra_deg, dec_deg):
    c = SkyCoord(ra=ra_deg * u.deg, dec=dec_deg * u.deg, frame='icrs')
    return (c.ra.to_string(unit=u.hour, sep=':', precision=2, pad=True),
            c.dec.to_string(unit=u.deg, sep=':', precision=1, pad=True))


def build_main_sheet(wb, frb_sched, twilight):
    if SHEET_NAME in wb.sheetnames:
        del wb[SHEET_NAME]
    ws = wb.create_sheet(SHEET_NAME)

    ws['A1'] = ('Lick 2026B-02 -- Night 2 (Sep 15/16): FRB repeats. '
                'FRB20260623C and FRB20260507D got redshifts night 1 -- '
                'NOT repeated.')
    ws['A1'].font = Font(size=12, bold=True)

    # Real DST-aware wall-clock label for this specific night (PDT in
    # September) -- see lick_2026b02_calculations.py's PST/ZoneInfo note.
    tz_label = twilight['evening_pst'].tzname()

    headers = ['Start (LST)', 'Start (UT)', f'Start ({tz_label})', 'Duration',
               f'End ({tz_label})', 'Action', 'target', 'RA', 'DEC', 'HA start',
               'HA end', 'mag(r)', 'origin', 'red side', 'blue side', 'Redshift;',
               'Comments']
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color='CCCCCC', end_color='CCCCCC', fill_type='solid')

    widths = [12, 12, 12, 12, 12, 30, 15, 12, 12, 10, 10, 8, 15, 15, 15, 10, 44]
    for col, w in zip('ABCDEFGHIJKLMNOPQ', widths):
        ws.column_dimensions[col].width = w

    current_row = 4

    def add_row(action, target='', ra='', dec='', mag_r='', origin='',
                red_side='', blue_side='', redshift='', comments='',
                start=None, end=None, ra_hours=None):
        nonlocal current_row
        lst_start = pst_to_lst(start)
        ut = pst_to_ut(start)
        duration_min = (end - start).total_seconds() / 60.0
        ha_start_str = ha_end_str = ''
        if ra_hours is not None:
            ha_start_str = f'{ha_at(ra_hours, lst_start):.2f}'
            ha_end_str = f'{ha_at(ra_hours, pst_to_lst(end)):.2f}'
        row = [f'{lst_start:.2f}', f'{ut:.2f}', start.strftime('%H:%M'),
               f'{duration_min:.0f} min' if duration_min > 0 else '',
               end.strftime('%H:%M') if duration_min > 0 else '',
               action, target, ra, dec, ha_start_str, ha_end_str, mag_r,
               origin, red_side, blue_side, redshift, comments]
        for col, val in enumerate(row, start=1):
            ws.cell(row=current_row, column=col, value=val)
        current_row += 1

    scheduled = frb_sched[frb_sched['scheduled']].copy()
    scheduled['start_pst'] = pd.to_datetime(scheduled['start_pst'], format='mixed')
    scheduled['end_pst'] = pd.to_datetime(scheduled['end_pst'], format='mixed')
    scheduled = scheduled.sort_values('start_pst').reset_index(drop=True)

    t = twilight['evening_pst']
    add_row('Focus', start=t, end=t + timedelta(minutes=5))
    t += timedelta(minutes=5)

    add_row('Standard Star', target=STD_STAR['name'], ra=STD_STAR['ra'],
             dec=STD_STAR['dec'], mag_r=STD_STAR['mag'],
             comments='Evening standard (same star used night 1 -- still '
                      'HA-legal at both ends of the night this date)',
             start=t, end=t + timedelta(minutes=5), ra_hours=STD_STAR['ra_hours'])
    t += timedelta(minutes=5)

    for _, targ in scheduled.iterrows():
        start = targ['start_pst'].to_pydatetime()
        end = targ['end_pst'].to_pydatetime()

        slew_start = max(t, start - timedelta(minutes=5))
        if slew_start > t:
            add_row('Idle -- no repeat target in HA-legal window',
                    comments='See night2_frb_repeat_analysis.csv for the full window table',
                    start=t, end=slew_start)
        add_row(f"slew to --> {targ['TNS']}", start=slew_start, end=start)

        ra_str, dec_str = format_ra_dec(targ['RA_deg'], targ['Dec_deg'])
        comment = (f"REPEAT ({targ['origin']}); "
                   + (f"P(O|x)={targ['POx']:.3f}; " if pd.notna(targ['POx']) else '')
                   + f"tags={targ['tags']}")
        if targ['skip_reason']:
            comment += f"; {targ['skip_reason']}"
        add_row('Science + overhead', target=targ['TNS'], ra=ra_str, dec=dec_str,
                mag_r=f"{targ['mag_r']:.2f}", origin=targ['origin'],
                red_side=targ['red_side'], blue_side=targ['blue_side'],
                comments=comment, start=start, end=end, ra_hours=targ['RA_hours'])
        t = end

    add_row('Standard Star', target=STD_STAR['name'], ra=STD_STAR['ra'],
             dec=STD_STAR['dec'], mag_r=STD_STAR['mag'],
             comments='Post-science standard',
             start=t + timedelta(minutes=5), end=t + timedelta(minutes=10),
             ra_hours=STD_STAR['ra_hours'])
    t += timedelta(minutes=10)
    add_row('Bias frames', start=t, end=t + timedelta(minutes=5))
    t += timedelta(minutes=5)
    add_row('Flat fields', start=t, end=t + timedelta(minutes=10))
    t += timedelta(minutes=10)

    margin_min = (twilight['morning_pst'] - t).total_seconds() / 60.0
    ws.cell(row=current_row + 1, column=1,
            value=f'Done by {t.strftime("%H:%M")} {tz_label}, {margin_min:.0f} min '
                  f'before 18-deg morning twilight '
                  f'({twilight["morning_pst"].strftime("%H:%M")} {tz_label}).')
    ws.cell(row=current_row + 1, column=1).font = Font(italic=True)

    ws.cell(row=3, column=19, value='Sidereal Time Calculator').font = Font(bold=True)
    ws.cell(row=4, column=19, value=tz_label).font = Font(bold=True)
    ws.cell(row=4, column=20, value='LST').font = Font(bold=True)
    for i, h in enumerate((0, 2, 4, 6, 8), start=5):
        st = twilight['evening_pst'] + timedelta(hours=h)
        ws.cell(row=i, column=19, value=st.strftime('%H:%M'))
        ws.cell(row=i, column=20, value=f'{pst_to_lst(st):.2f}')

    return ws


def build_spare_sheet(wb, spare_sched, spare_analysis, twilight):
    if SPARE_SHEET_NAME in wb.sheetnames:
        del wb[SPARE_SHEET_NAME]
    ws = wb.create_sheet(SPARE_SHEET_NAME)

    ws['A1'] = ('Lick 2026B-02 -- Spare/filler targets, author-supplied '
                '2026-09-15 (non-FRB; RA/Dec/mag given directly, filter '
                'varies -- used as an r-mag proxy for the Kast exposure '
                'heuristic, not a science-grade calibration).')
    ws['A1'].font = Font(size=12, bold=True)

    tz_label = twilight['evening_pst'].tzname()
    headers = [f'Start ({tz_label})', f'End ({tz_label})', 'target', 'RA (h)',
               'Dec (deg)', 'given mag', 'filter', 'red side', 'blue side',
               'Comments']
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color='D6E4F0', end_color='D6E4F0', fill_type='solid')
    widths = [12, 12, 12, 10, 10, 10, 8, 10, 10, 60]
    for col, w in zip('ABCDEFGHIJ', widths):
        ws.column_dimensions[col].width = w

    filt_lookup = {'2026yly': 'C', '2026sqf': 'V', '2026abgb': 'g',
                   '2026yum': 'r', '2026zyg': 'L', '2025pht': 'r'}

    row = 4
    scheduled = spare_sched[spare_sched['scheduled']].copy() if len(spare_sched) else spare_sched
    if len(scheduled):
        scheduled['start_pst'] = pd.to_datetime(scheduled['start_pst'], format='mixed')
        scheduled['end_pst'] = pd.to_datetime(scheduled['end_pst'], format='mixed')
        scheduled = scheduled.sort_values('start_pst').reset_index(drop=True)
        for _, r in scheduled.iterrows():
            ws.cell(row=row, column=1, value=r['start_pst'].strftime('%H:%M'))
            ws.cell(row=row, column=2, value=r['end_pst'].strftime('%H:%M'))
            ws.cell(row=row, column=3, value=r['TNS'])
            ws.cell(row=row, column=4, value=f"{r['RA_hours']:.3f}")
            ws.cell(row=row, column=5, value=f"{r['Dec_deg']:.3f}")
            ws.cell(row=row, column=6,
                    value=f"{r['mag_r']:.1f}" if pd.notna(r['mag_r']) else 'n/a')
            ws.cell(row=row, column=7, value=filt_lookup.get(r['TNS'], ''))
            ws.cell(row=row, column=8, value=r['red_side'])
            ws.cell(row=row, column=9, value=r['blue_side'])
            ws.cell(row=row, column=10, value=r.get('skip_reason', '') or '')
            row += 1

    row += 1
    ws.cell(row=row, column=1, value='Not scheduled / not HA-legal tonight:').font = Font(bold=True)
    row += 1
    scheduled_names = set(scheduled['TNS']) if len(scheduled) else set()
    for _, r in spare_analysis.iterrows():
        if r['TNS'] in scheduled_names:
            continue
        if r['observable']:
            note = 'HA-legal but no idle gap large enough tonight'
        else:
            entry_time, exit_time, sunset, sunrise = never_legal_diagnostic(
                r['TNS'], r['RA_hours'], twilight)
            overlaps = entry_time < sunrise and exit_time > sunset
            note = (f"Shane HA window would run {entry_time:%H:%M}->{exit_time:%m-%d %H:%M} "
                    f"{exit_time.tzname()} (sunrise {sunrise:%H:%M}) -- "
                    + ("marginally overlaps pre-sunrise twilight, could be "
                       "pushed as a dawn attempt if the author wants to go "
                       "past the formal 18-deg cutoff"
                       if overlaps else
                       "does not open until after sunrise -- not reachable "
                       "tonight under the real Shane HA mount limit"))
        ws.cell(row=row, column=3, value=r['TNS'])
        ws.cell(row=row, column=10, value=note)
        row += 1

    return ws


def main():
    twilight = find_twilight(DATE)
    frb_sched = pd.read_csv(f'{OUT_DIR}/night2_frb_repeat_schedule.csv')
    spare_sched = pd.read_csv(f'{OUT_DIR}/night2_spare_schedule.csv')
    spare_analysis = pd.read_csv(f'{OUT_DIR}/night2_spare_analysis.csv')

    wb = load_workbook(WB_PATH)
    build_main_sheet(wb, frb_sched, twilight)
    build_spare_sheet(wb, spare_sched, spare_analysis, twilight)
    wb.save(WB_PATH)
    print(f'Added "{SHEET_NAME}" and "{SPARE_SHEET_NAME}" sheets to {WB_PATH}')
    print(f'Sheets now: {wb.sheetnames}')


if __name__ == '__main__':
    main()
